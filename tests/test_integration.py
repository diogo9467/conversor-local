"""Exercise the shipped C# app AND frozen engines, including publication beside sources."""
import concurrent.futures
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import uuid
import zipfile
import pytest
import pymupdf as fitz
from docx import Document
from pptx import Presentation
from make_fixtures import make,FIXTURES,ROOT

APP=Path(os.environ.get('CONVERSOR_TEST_APP',ROOT/'.build'/'app'/'ConversorLocal.exe'))
REPORTS=ROOT/'test-results'/'reports'

@pytest.fixture(scope='session',autouse=True)
def setup():
    make();REPORTS.mkdir(parents=True,exist_ok=True)
    assert APP.exists()

def convert(files,target,ocr=True):
    report=REPORTS/(uuid.uuid4().hex+'.json')
    cmd=[str(APP),'--no-ui','--to',target,'--ocr' if ocr else '--no-ocr','--report',str(report),'--',*map(str,files)]
    before={str(p):hashlib.sha256(Path(p).read_bytes()).hexdigest() for p in files}
    run=subprocess.run(cmd,capture_output=True,timeout=300)
    assert report.exists(),(run.returncode,run.stderr,run.stdout)
    results=json.loads(report.read_text(encoding='utf-8-sig'))
    for source in files:assert hashlib.sha256(source.read_bytes()).hexdigest()==before[str(source)]
    for item in results:
        if item['Success']:
            output=Path(item['Output']);assert output.parent==Path(item['Source']).parent
            assert output.exists() and output.stat().st_size>0
    return run.returncode,results

A=FIXTURES/'Pasta A com acentos';B=FIXTURES/'Outra pasta'

def test_docx_pdf_multipage_header_table_image():
    code,results=convert([A/'Artigo científico & revisão.docx'],'pdf')
    assert code==0,results
    with fitz.open(results[0]['Output']) as pdf:
        assert len(pdf)==3
        assert all(f'MARCADOR_DOCX_{i+1}' in pdf[i].get_text() for i in range(3))
        assert all('CABECALHO' in p.get_text() and p.get_images() for p in pdf)
        assert 'Quantidade' in pdf[0].get_text()

def test_docx_pdf_external_template_keeps_embedded_content():
    from test_external_templates import with_template
    source=with_template(A/'Artigo científico & revisão.docx',A/'Documento com modelo externo.docx')
    code,results=convert([source],'pdf');assert code==0,results
    with fitz.open(results[0]['Output']) as pdf:
        assert len(pdf)==3
        assert all(f'MARCADOR_DOCX_{i+1}' in pdf[i].get_text() for i in range(3))
        assert all('CABECALHO' in p.get_text() and p.get_images() for p in pdf)
        assert 'Quantidade' in pdf[0].get_text()

def test_docx_pdf_does_not_fetch_external_template():
    import http.server
    import threading
    from test_external_templates import with_template
    hits=[]
    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path);self.send_response(404);self.end_headers()
        def do_HEAD(self):self.do_GET()
        def do_POST(self):self.do_GET()
        def log_message(self,*args):pass
    with http.server.ThreadingHTTPServer(('127.0.0.1',0),Handler) as server:
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        try:
            target=f'http://127.0.0.1:{server.server_port}/external-template.dotm'
            source=with_template(A/'Artigo científico & revisão.docx',A/'Modelo HTTP não deve ser consultado.docx',target=target)
            code,results=convert([source],'pdf');assert code==0,results
            assert hits==[],f'External template was requested: {hits}'
        finally:server.shutdown();thread.join(timeout=5)

def test_docx_pdf_genuine_external_image_still_fails():
    from test_external_templates import with_template
    source=with_template(A/'Artigo científico & revisão.docx',A/'Documento com imagem realmente externa.docx',extra_type='image')
    code,results=convert([source],'pdf')
    assert code==1 and results[0]['Output'] is None
    assert 'imagem vinculada' in results[0]['Error']

def test_pptx_pdf_all_slides_and_aspect():
    code,results=convert([B/'Seminário com imagens.pptx'],'pdf');assert code==0,results
    with fitz.open(results[0]['Output']) as pdf:
        assert len(pdf)==3
        assert all(f'MARCADOR_PPTX_{i+1}' in pdf[i].get_text() for i in range(3))
        assert abs(pdf[0].rect.width/pdf[0].rect.height-16/9)<.01
        assert pdf[0].get_images()

def test_pdf_docx_editable_text_tables_images():
    code,results=convert([A/'Pesquisa com texto e tabelas.pdf'],'docx');assert code==0,results
    doc=Document(results[0]['Output'])
    text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert all(f'MARCADOR_PDF_{i+1}' in text for i in range(3))
    assert len(doc.tables)>=1
    with zipfile.ZipFile(results[0]['Output']) as z:assert any(n.startswith('word/media/') for n in z.namelist())

def test_pdf_pptx_every_page_embedded_image():
    code,results=convert([A/'Pesquisa com texto e tabelas.pdf'],'pptx');assert code==0,results
    deck=Presentation(results[0]['Output']);assert len(deck.slides)==3
    assert all(len(slide.shapes)==1 and slide.shapes[0].image.size[0]>1000 for slide in deck.slides)

def test_scan_with_ocr_editable():
    code,results=convert([FIXTURES/'Digitalizados'/'Digitalização português.pdf'],'docx');assert code==0,results
    doc=Document(results[0]['Output'])
    text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert 'MARCADOR' in text and 'Pesquisa' in text
    assert any('OCR' in w for w in results[0]['Warnings'])

def test_scan_without_ocr_clear_failure():
    code,results=convert([FIXTURES/'Digitalizados'/'Digitalização português.pdf'],'docx',ocr=False)
    assert code==1 and 'OCR' in results[0]['Error'] and results[0]['Output'] is None

def test_batch_isolates_bad_file():
    code,results=convert([B/'Arquivo corrompido.pdf',A/'Pesquisa com texto e tabelas.pdf'],'pptx')
    assert code==1 and not results[0]['Success'] and results[1]['Success']

def test_password_protection_clear_failure():
    code,results=convert([B/'Protegido por senha.pdf'],'docx');assert code==1
    assert 'senha' in results[0]['Error'] and results[0]['Output'] is None

def test_bad_office_clear_failure():
    code,results=convert([B/'Office corrompido.docx'],'pdf');assert code==1
    assert not results[0]['Success'] and results[0]['Output'] is None

def test_existing_target_preserved():
    source=B/'Colisão de nome.pdf';shutil.copy2(A/'Pesquisa com texto e tabelas.pdf',source)
    target=source.with_suffix('.pptx');target.write_bytes(b'KEEP THIS EXISTING FILE')
    code,results=convert([source],'pptx');assert code==0,results
    assert target.read_bytes()==b'KEEP THIS EXISTING FILE'
    assert Path(results[0]['Output']).stem.startswith('Colisão de nome (')

def test_simultaneous_same_source_unique_outputs():
    source=A/'Pesquisa com texto e tabelas.pdf'
    with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
        futures=[pool.submit(convert,[source],'pptx') for _ in range(2)]
        results=[f.result() for f in futures]
    assert all(code==0 for code,_ in results),results
    assert len({r[0]['Output'] for _,r in results})==2

def test_mixed_page_sizes_no_crop():
    code,results=convert([B/'Páginas de proporções diferentes.pdf'],'pptx');assert code==0,results
    deck=Presentation(results[0]['Output'])
    assert len(deck.slides)==3
    assert all(s.shapes[0].left>=0 and s.shapes[0].top>=0 for s in deck.slides)
    assert any('diferentes' in w for w in results[0]['Warnings'])

def test_long_path_and_special_characters():
    folder=FIXTURES/'long-path'/('a'*90)/('b'*90)/('c'*60);folder.mkdir(parents=True,exist_ok=True)
    source=folder/'Ação $ & [revisão] #1.pdf';shutil.copy2(A/'Pesquisa com texto e tabelas.pdf',source)
    assert len(str(source))>260
    code,results=convert([source],'pptx');assert code==0,results

def test_no_partial_outputs():
    assert not list(FIXTURES.rglob('*.partial'))
    assert not list(FIXTURES.rglob('*.probe'))
