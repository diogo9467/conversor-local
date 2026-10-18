"""Generated, deterministic local test documents; no personal documents are used."""
from pathlib import Path
import io
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
import pymupdf as fitz

ROOT=Path(__file__).resolve().parents[1]
FIXTURES=ROOT/'test-results'/'fixtures'

def make():
    for folder in ['Pasta A com acentos','Outra pasta','Digitalizados']:(FIXTURES/folder).mkdir(parents=True,exist_ok=True)
    a=FIXTURES/'Pasta A com acentos'; b=FIXTURES/'Outra pasta'; scan=FIXTURES/'Digitalizados'
    image=Image.new('RGB',(700,350),'#E8F3F9');d=ImageDraw.Draw(image)
    d.rectangle((30,35,220,310),fill='#163D61');d.ellipse((290,50,550,310),fill='#259B9A');d.line((150,180,420,180,620,80),fill='#E47732',width=14)
    image.save(a/'figura.png')
    doc=Document()
    section=doc.sections[0];section.header.paragraphs[0].text='CABECALHO · Conversor Local'
    section.footer.paragraphs[0].text='Rodape de verificacao · documento de teste'
    for page in range(3):
        if page:doc.add_page_break()
        doc.add_heading(f'Revisão científica - página {page+1}',0)
        doc.add_paragraph('Texto editável com acentuação: ação, coração, português. MARCADOR_DOCX_'+str(page+1))
        table=doc.add_table(rows=1,cols=3);table.style='Table Grid'
        for c,v in zip(table.rows[0].cells,['Item','Quantidade','Valor']):c.text=v
        for j in range(3):
            for c,v in zip(table.add_row().cells,[f'Linha {j+1}',str(j+2),f'R$ {j*10+15},00']):c.text=v
        doc.add_picture(str(a/'figura.png'),width=DocInches(5.5))
        doc.add_paragraph('Parágrafo final para conferir margens, organização e paginação.')
    doc.save(a/'Artigo científico & revisão.docx')

    deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
    for index in range(3):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        if index==2: slide._element.set('show','0')  # Hidden slide must still be exported.
        title=slide.shapes.add_textbox(Inches(.6),Inches(.4),Inches(12),Inches(1))
        p=title.text_frame.paragraphs[0];p.text=f'Seminário - slide {index+1}';p.font.size=Pt(32);p.font.color.rgb=RGBColor.from_string('163D61')
        text=slide.shapes.add_textbox(Inches(.6),Inches(1.6),Inches(5),Inches(4))
        text.text_frame.text=f'MARCADOR_PPTX_{index+1}\nTexto com ação e coração\nLayout {index+1}: texto e imagem'
        if index!=1:slide.shapes.add_picture(str(a/'figura.png'),Inches(6.1),Inches(2),width=Inches(6.4))
        else:
            table=slide.shapes.add_table(3,3,Inches(6.1),Inches(2),Inches(6.4),Inches(2)).table
            for r in range(3):
                for c in range(3):table.cell(r,c).text=f'Dado {r+1}/{c+1}'
    deck.save(b/'Seminário com imagens.pptx')

    pdf=a/'Pesquisa com texto e tabelas.pdf'
    cv=canvas.Canvas(str(pdf),pagesize=A4)
    for page in range(3):
        cv.setFont('Helvetica-Bold',22);cv.drawString(50,780,f'Pesquisa - pagina {page+1}')
        cv.setFont('Helvetica',13);cv.drawString(50,744,f'MARCADOR_PDF_{page+1} Texto selecionavel e editavel.')
        cv.drawString(50,720,'Dados de exemplo para conversao local offline.')
        for row in range(5):
            y=660-row*28
            cv.line(50,y,530,y)
            cv.drawString(60,y-19,f'Item {row+1}');cv.drawString(230,y-19,str(row*15));cv.drawString(380,y-19,f'Valor {row+1}')
        for x in [50,210,360,530]:cv.line(x,660,x,520)
        cv.line(50,520,530,520)
        cv.drawImage(str(a/'figura.png'),50,275,width=480,height=240)
        cv.setFont('Helvetica',10);cv.drawString(50,60,f'Rodape - {page+1} de 3')
        cv.showPage()
    cv.save()
    with fitz.open(pdf) as native:
        scanned=fitz.open()
        for page in native:
            pix=page.get_pixmap(dpi=150)
            output=scanned.new_page(width=page.rect.width,height=page.rect.height)
            output.insert_image(output.rect,stream=pix.tobytes('png'))
        scanned.save(scan/'Digitalização português.pdf');scanned.close()
        native.save(b/'Protegido por senha.pdf',encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='test-only',owner_pw='test-only-owner')
    mixed=fitz.open()
    for w,h in [(400,600),(800,400),(400,600)]:
        page=mixed.new_page(width=w,height=h);page.insert_text((30,60),f'Pagina {len(mixed)} com proporcao {w}/{h}')
    mixed.save(b/'Páginas de proporções diferentes.pdf');mixed.close()
    (b/'Arquivo corrompido.pdf').write_bytes(b'%PDF-1.7\nNOT A VALID PDF\n')
    (b/'Office corrompido.docx').write_bytes(b'broken zip file')
    return FIXTURES

if __name__=='__main__':print(make())
