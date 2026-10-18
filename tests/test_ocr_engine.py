"""Fast regression checks for page-specific OCR reconstruction before freezing."""
import importlib.util
from pathlib import Path
import sys
import pymupdf as fitz
from docx import Document
from make_fixtures import make,FIXTURES,ROOT

def test_mixed_scanned_and_native_pdf(tmp_path):
    sys.path.insert(0,str(ROOT/'engine'))
    from worker import pdf_docx,validate_ooxml
    make()
    native=fitz.open(FIXTURES/'Pasta A com acentos'/'Pesquisa com texto e tabelas.pdf')
    scan=fitz.open(FIXTURES/'Digitalizados'/'Digitalização português.pdf')
    mixed=fitz.open();mixed.insert_pdf(native,from_page=0,to_page=0);mixed.insert_pdf(scan,from_page=0,to_page=0)
    source=tmp_path/'mixed.pdf';mixed.save(source);mixed.close();native.close();scan.close()
    output=tmp_path/'mixed.docx'
    assert pdf_docx(source,output,ROOT/'.build'/'app',tmp_path,True)==2
    validate_ooxml(output,'docx')
    doc=Document(output)
    text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert text.count('MARCADOR')>=2
    assert len(doc.inline_shapes)>=1
