"""Regression: optional Word templates must not block or contact external locations."""
import hashlib
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from xml.dom import minidom
import zipfile
import pytest
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
from worker import ConversionError,prepare_office_input,validate_ooxml

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'

def with_template(source,dest,target='http://127.0.0.1:1/template.dotm',extra_type=None):
    """Append relationships to a real DOCX without changing its content or media."""
    with zipfile.ZipFile(source) as archive:
        settings=minidom.parseString(archive.read('word/settings.xml'))
        node=settings.createElementNS(W,'w:attachedTemplate');node.setAttributeNS(R,'r:id','rIdRegressionTemplate')
        settings.documentElement.appendChild(node)
        link=settings.createElementNS(W,'w:linkStyles');settings.documentElement.appendChild(link)
        rels=ET.Element('{'+P+'}Relationships')
        ET.SubElement(rels,'{'+P+'}Relationship',Id='rIdRegressionTemplate',Type=R+'/attachedTemplate',Target=target,TargetMode='External')
        if extra_type:ET.SubElement(rels,'{'+P+'}Relationship',Id='rIdExternalOther',Type=R+'/'+extra_type,Target='http://127.0.0.1:1/resource',TargetMode='External')
        replacement={'word/settings.xml':settings.toxml(encoding='utf-8'),'word/_rels/settings.xml.rels':ET.tostring(rels,encoding='utf-8',xml_declaration=True)}
        with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as output:
            for entry in archive.infolist():output.writestr(entry,replacement.pop(entry.filename,archive.read(entry)))
            for name,data in replacement.items():output.writestr(name,data)
    return dest

@pytest.fixture
def docx(tmp_path):
    path=tmp_path/'documento.docx';doc=Document();doc.add_heading('Texto preservado',0)
    doc.add_paragraph('MARCADOR_TEMPLATE: ação, coração e revisão.');doc.add_table(rows=2,cols=2).cell(0,0).text='Tabela incorporada'
    doc.save(path)
    return path

def test_detaches_template_only_and_preserves_original_and_parts(docx,tmp_path):
    source=with_template(docx,tmp_path/'origem com acentos.docx')
    before=hashlib.sha256(source.read_bytes()).hexdigest();job=tmp_path/'job';job.mkdir()
    with pytest.raises(ConversionError):validate_ooxml(source,'docx',check_external=True)
    prepared=prepare_office_input(source,job)
    assert prepared!=source and hashlib.sha256(source.read_bytes()).hexdigest()==before
    validate_ooxml(prepared,'docx',check_external=True)
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(prepared) as sanitized:
        assert set(original.namelist())==set(sanitized.namelist())
        for name in original.namelist():
            if name not in ['word/settings.xml','word/_rels/settings.xml.rels']:assert original.read(name)==sanitized.read(name),name
        assert b'attachedTemplate' not in sanitized.read('word/settings.xml')
        assert b'linkStyles' not in sanitized.read('word/settings.xml')
        xml=minidom.parseString(sanitized.read('word/settings.xml'))
        for prefix in xml.documentElement.getAttributeNS('http://schemas.openxmlformats.org/markup-compatibility/2006','Ignorable').split():
            assert xml.documentElement.hasAttribute('xmlns:'+prefix)
    doc=Document(prepared)
    assert 'MARCADOR_TEMPLATE' in '\n'.join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(0,0).text=='Tabela incorporada'

@pytest.mark.parametrize('target',['file:///C:/old/template.dot','https://example.invalid/template.dotm',r'\\server.invalid\share\template.dotx'])
def test_does_not_resolve_template_target(docx,tmp_path,target):
    source=with_template(docx,tmp_path/'source.docx',target=target);job=tmp_path/'job';job.mkdir()
    assert prepare_office_input(source,job).exists()

@pytest.mark.parametrize('kind',['image','oleObject','aFChunk'])
def test_genuine_external_content_still_rejected(docx,tmp_path,kind):
    source=with_template(docx,tmp_path/'source.docx',extra_type=kind);job=tmp_path/'job';job.mkdir()
    with pytest.raises(ConversionError):prepare_office_input(source,job)

def test_self_contained_docx_does_not_need_rewrite(docx,tmp_path):
    job=tmp_path/'job';job.mkdir()
    assert prepare_office_input(docx,job)==docx
    assert list(job.iterdir())==[]
