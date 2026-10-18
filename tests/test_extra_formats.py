"""Exercise every catalog route through the shipped host and frozen worker."""
import io
import json
from pathlib import Path
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET

import pytest
import pymupdf as fitz
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from make_fixtures import ROOT, FIXTURES, make
from test_integration import convert

sys.path.insert(0,str(ROOT/'engine'))
from extra_formats import ConversionError, validate_odf, validate_xlsx, validate_rtf, read_text

CATALOG=json.loads((ROOT/'formats.json').read_text(encoding='utf-8'))
EXTRA=FIXTURES/'Novos formatos'
NS='''xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 xmlns:xlink="http://www.w3.org/1999/xlink"
 xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"
 xmlns:of="urn:oasis:names:tc:opendocument:xmlns:of:1.2"'''
MARKER='MARCADOR_NOVO ação coração'


def odf(path,kind,body,extra=None):
    mime='application/vnd.oasis.opendocument.'+{'odt':'text','odp':'presentation','ods':'spreadsheet'}[kind]
    content=f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content {NS} office:version="1.3"><office:body>{body}</office:body></office:document-content>'''
    styles=f'''<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles {NS} office:version="1.3"><office:styles/>
<office:automatic-styles><style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="28cm" fo:page-height="15.75cm" style:print-orientation="landscape"/></style:page-layout></office:automatic-styles>
<office:master-styles><style:master-page style:name="Default" style:page-layout-name="pm1"/></office:master-styles></office:document-styles>'''
    entries={'content.xml':content.encode(),'styles.xml':styles.encode(),**(extra or {})}
    manifest=f'''<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3"><manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'''
    for name in entries:
        media='text/xml' if name.endswith('.xml') else 'image/png'
        manifest+=f'<manifest:file-entry manifest:full-path="{name}" manifest:media-type="{media}"/>'
    manifest+='</manifest:manifest>'
    with zipfile.ZipFile(path,'w') as archive:
        archive.writestr('mimetype',mime,compress_type=zipfile.ZIP_STORED)
        for name,data in entries.items():archive.writestr(name,data,compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr('META-INF/manifest.xml',manifest)


def xlsx(path):
    main='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    types='http://schemas.openxmlformats.org/package/2006/content-types'
    parts={
        '[Content_Types].xml':f'<Types xmlns="{types}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'+''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in (1,2))+'</Types>',
        '_rels/.rels':f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        'xl/workbook.xml':f'<workbook xmlns="{main}" xmlns:r="{rel}"><sheets><sheet name="Pesquisa" sheetId="1" r:id="rId1"/><sheet name="Resumo" sheetId="2" r:id="rId2"/></sheets><calcPr calcId="191029" fullCalcOnLoad="1"/></workbook>',
        'xl/_rels/workbook.xml.rels':'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'+''.join(f'<Relationship Id="rId{i}" Type="{rel}/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in (1,2))+'</Relationships>'
    }
    for i in (1,2):
        parts[f'xl/worksheets/sheet{i}.xml']=f'''<worksheet xmlns="{main}"><sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>
<cols><col min="1" max="1" width="50" customWidth="1"/><col min="2" max="2" width="15" customWidth="1"/></cols><sheetData>
<row r="1"><c r="A1" t="inlineStr"><is><t>MARCADOR_ABA_{i} ação coração</t></is></c></row>
<row r="2"><c r="A2" t="inlineStr"><is><t>Item</t></is></c><c r="B2" t="inlineStr"><is><t>Valor</t></is></c></row>
<row r="3"><c r="A3" t="inlineStr"><is><t>Pesquisa</t></is></c><c r="B3"><v>12</v></c></row>
<row r="4"><c r="A4" t="inlineStr"><is><t>Revisão</t></is></c><c r="B4"><v>17</v></c></row>
<row r="5"><c r="A5" t="inlineStr"><is><t>Total</t></is></c><c r="B5"><f>SUM(B3:B4)</f><v>29</v></c></row>
</sheetData><pageMargins left="0.5" right="0.5" top="0.5" bottom="0.5" header="0.2" footer="0.2"/><pageSetup paperSize="9" fitToWidth="1" fitToHeight="1"/></worksheet>'''
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in parts.items():archive.writestr(name,data.encode())


@pytest.fixture(scope='session',autouse=True)
def extended_fixtures():
    EXTRA.mkdir(parents=True,exist_ok=True)
    make()
    doc=Document();doc.add_heading(MARKER,0);doc.add_paragraph('Texto editável, 12 + 17 = 29.')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='Valor';table.cell(1,0).text='Pesquisa';table.cell(1,1).text='29'
    doc.add_picture(str(FIXTURES/'Pasta A com acentos'/'figura.png'),width=Inches(3));doc.save(EXTRA/'origem.docx')
    (EXTRA/'origem.txt').write_text(MARKER+'\nTexto editável, 12 + 17 = 29.\n',encoding='utf-8')
    (EXTRA/'origem.rtf').write_bytes(rb'{\rtf1\ansi\ansicpg1252\deff0 {\fonttbl{\f0 Arial;}}\f0\fs24 MARCADOR_NOVO a\u231?\u227?o cora\u231?\u227?o\par Texto edit\u225?vel, 12 + 17 = 29.}')
    picture=(FIXTURES/'Pasta A com acentos'/'figura.png').read_bytes()
    odf(EXTRA/'origem.odt','odt',f'''<office:text><text:h text:outline-level="1">{MARKER}</text:h><text:p>Texto editável, 12 + 17 = 29.</text:p><table:table table:name="Dados"><table:table-column table:number-columns-repeated="2"/><table:table-row><table:table-cell office:value-type="string"><text:p>Pesquisa</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="29"><text:p>29</text:p></table:table-cell></table:table-row></table:table><text:p><draw:frame draw:name="Figura" text:anchor-type="as-char" svg:width="8cm" svg:height="4cm"><draw:image xlink:href="Pictures/figura.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></text:p></office:text>''',{'Pictures/figura.png':picture})
    slides=''
    for i in (1,2):
        slides+=f'''<draw:page draw:name="Slide {i}" draw:master-page-name="Default"><draw:frame svg:x="1cm" svg:y="1cm" svg:width="25cm" svg:height="2cm"><draw:text-box><text:p>MARCADOR_SLIDE_{i} ação coração</text:p></draw:text-box></draw:frame><draw:frame svg:x="2cm" svg:y="4cm" svg:width="16cm" svg:height="8cm"><draw:image xlink:href="Pictures/figura.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame></draw:page>'''
    odf(EXTRA/'origem.odp','odp','<office:presentation>'+slides+'</office:presentation>',{'Pictures/figura.png':picture})
    deck=Presentation();deck.slide_width=12192000;deck.slide_height=6858000
    for i in (1,2):
        slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(100000,100000,8000000,1000000).text=f'MARCADOR_SLIDE_{i} ação coração'
        slide.shapes.add_picture(io.BytesIO(picture),100000,1500000,width=6000000)
    deck.save(EXTRA/'origem.pptx')
    sheets=''
    for i in (1,2):
        sheets+=f'''<table:table table:name="Aba {i}"><table:table-column table:number-columns-repeated="2"/>
<table:table-row><table:table-cell office:value-type="string"><text:p>MARCADOR_ABA_{i} ação coração</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>Item</text:p></table:table-cell><table:table-cell office:value-type="string"><text:p>Valor</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>Pesquisa</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="12"><text:p>12</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>Revisão</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="17"><text:p>17</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>Total</text:p></table:table-cell><table:table-cell table:formula="of:=SUM([.B3:.B4])" office:value-type="float" office:value="29"><text:p>29</text:p></table:table-cell></table:table-row></table:table>'''
    odf(EXTRA/'origem.ods','ods','<office:spreadsheet>'+sheets+'</office:spreadsheet>')
    xlsx(EXTRA/'origem.xlsx')
    shutil.copy2(FIXTURES/'Pasta A com acentos'/'Pesquisa com texto e tabelas.pdf',EXTRA/'origem.pdf')
    image=Image.new('RGBA',(360,240),(0,0,0,0));draw=ImageDraw.Draw(image);draw.rectangle((20,20,180,200),fill='#163D61');draw.ellipse((150,40,330,220),fill='#25B5AC')
    for ext,fmt in [('png','PNG'),('webp','WEBP'),('bmp','BMP'),('jpg','JPEG'),('jpeg','JPEG'),('tif','TIFF'),('tiff','TIFF')]:
        image.convert('RGB').save(EXTRA/('origem.'+ext),format=fmt) if ext in ('bmp','jpg','jpeg') else image.save(EXTRA/('origem.'+ext),format=fmt,**({'lossless':True} if ext=='webp' else {}))
    frames=[Image.new('RGB',(200,300),'#163D61'),Image.new('RGB',(400,200),'#25B5AC')]
    frames[0].save(EXTRA/'multiplas.tiff',save_all=True,append_images=frames[1:],compression='tiff_lzw')
    frames[0].resize((200,300)).save(EXTRA/'animada.webp',save_all=True,append_images=[frames[1].resize((200,300))],duration=100,loop=0,lossless=True)


ROUTES=[(ext,target) for ext,source in CATALOG['sources'].items() for target in source['targets']]


def package_text(path):
    with zipfile.ZipFile(path) as archive:
        return '\n'.join(''.join(ET.fromstring(archive.read(name)).itertext()) for name in archive.namelist() if name.endswith('.xml') and ('content' in name or '/slides/slide' in name or '/worksheets/' in name or name=='xl/sharedStrings.xml'))


@pytest.mark.parametrize('ext,target',ROUTES,ids=[f'{e}-to-{t}' for e,t in ROUTES])
def test_every_supported_route(ext,target):
    code,results=convert([EXTRA/('origem.'+ext)],target);assert code==0,results
    output=Path(results[0]['Output']);family=CATALOG['sources'][ext]['family']
    assert output.suffix=='.'+target
    if target=='pdf':
        with fitz.open(output) as pdf:
            text=''.join(page.get_text() for page in pdf)
            if family=='word':assert 'MARCADOR_NOVO' in text and 'ação coração' in text and '29' in text
            if family=='slides':assert len(pdf)==2 and all(f'MARCADOR_SLIDE_{i+1}' in pdf[i].get_text() and pdf[i].get_images() for i in range(2))
            if family=='sheet':assert all(f'MARCADOR_ABA_{i}' in text for i in (1,2)) and '29' in text
            if family=='image':assert len(pdf)==1 and pdf[0].get_images()
    elif target=='docx':
        document=Document(output);text='\n'.join([p.text for p in document.paragraphs]+[c.text for t in document.tables for r in t.rows for c in r.cells])
        assert ('MARCADOR_PDF_1' if ext=='pdf' else 'MARCADOR_NOVO') in text
        if ext!='pdf':assert 'ação coração' in text and '29' in text
        if ext=='odt':assert document.tables
    elif target=='txt':
        text=output.read_text(encoding='utf-8-sig');assert ('MARCADOR_PDF_1' if ext=='pdf' else 'MARCADOR_NOVO') in text
        if ext!='pdf':assert 'ação coração' in text and '29' in text
    elif target=='rtf':assert b'MARCADOR_NOVO' in output.read_bytes() and b'29' in output.read_bytes()
    elif target in ('odt','odp','ods','xlsx'):
        text=package_text(output)
        if family=='word':assert 'MARCADOR_NOVO' in text and 'ação coração' in text and '29' in text
        if family=='slides':assert all(f'MARCADOR_SLIDE_{i}' in text for i in (1,2))
        if family=='sheet':
            assert all(f'MARCADOR_ABA_{i}' in text for i in (1,2)) and '29' in text
            with zipfile.ZipFile(output) as archive:
                if target=='ods':assert b'SUM' in archive.read('content.xml')
                else:assert any(b'SUM' in archive.read(n) for n in archive.namelist() if n.startswith('xl/worksheets/'))
    elif target=='pptx':
        deck=Presentation(output);assert len(deck.slides)==(3 if ext=='pdf' else 2)
        if ext=='odp':assert all(f'MARCADOR_SLIDE_{i+1}' in '\n'.join(s.text for s in deck.slides[i].shapes if s.has_text_frame) for i in range(2))
    else:
        with Image.open(output) as image:
            assert image.size==(360,240) and getattr(image,'n_frames',1)==1
            image.load()
            if ext=='png' and target in ('webp','tiff'):assert image.getpixel((0,0))[-1]==0
            if ext=='png' and target in ('jpg','bmp'):assert all(v>=250 for v in image.convert('RGB').getpixel((0,0)))
    with (ROOT/'test-results'/'extended-routes.jsonl').open('a',encoding='utf-8') as stream:
        stream.write(json.dumps({'source':str(EXTRA/('origem.'+ext)),'target':target,'output':str(output)},ensure_ascii=False)+'\n')


@pytest.mark.parametrize('name,target,count',[('multiplas.tiff','pdf',2),('animada.webp','pdf',2),('animada.webp','tiff',2)])
def test_all_image_frames_preserved(name,target,count):
    code,results=convert([EXTRA/name],target);assert code==0,results
    if target=='pdf':
        with fitz.open(results[0]['Output']) as pdf:assert len(pdf)==count and all(page.get_images() for page in pdf)
    else:
        with Image.open(results[0]['Output']) as image:assert image.n_frames==count


@pytest.mark.parametrize('name',['multiplas.tiff','animada.webp'])
def test_multiple_frames_never_silently_discarded(name):
    code,results=convert([EXTRA/name],'png')
    assert code==1 and results[0]['Output'] is None and 'várias páginas' in results[0]['Error']


def test_exif_orientation_preserves_displayed_dimensions():
    source=EXTRA/'orientacao.jpeg';exif=Image.Exif();exif[274]=6
    Image.new('RGB',(100,200),'#163D61').save(source,exif=exif)
    code,results=convert([source],'png');assert code==0,results
    with Image.open(results[0]['Output']) as image:assert image.size==(200,100) and image.getexif().get(274,1)==1


@pytest.mark.parametrize('encoding',['utf-8','utf-16','cp1252'])
def test_txt_encodings(encoding):
    source=EXTRA/f'codificacao-{encoding}.txt';source.write_text(MARKER,encoding=encoding)
    code,results=convert([source],'docx');assert code==0,results
    assert MARKER in '\n'.join(p.text for p in Document(results[0]['Output']).paragraphs)


def test_scanned_pdf_to_text_with_and_without_ocr():
    source=FIXTURES/'Digitalizados'/'Digitalização português.pdf'
    code,results=convert([source],'txt',ocr=False);assert code==1 and results[0]['Output'] is None and 'OCR' in results[0]['Error']
    code,results=convert([source],'txt',ocr=True);assert code==0,results
    assert 'MARCADOR' in Path(results[0]['Output']).read_text(encoding='utf-8-sig') and any('OCR' in w for w in results[0]['Warnings'])


def rewrite_zip(source,output,parts):
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as result:
        for name in archive.namelist():result.writestr(name,parts.pop(name,archive.read(name)))
        for name,data in parts.items():result.writestr(name,data)


@pytest.mark.parametrize('kind',['xlsx-link','xlsx-formula','ods-link','odt-image','rtf-field','bad-image','empty-text','binary-text'])
def test_invalid_or_external_content_publishes_nothing(kind):
    if kind=='xlsx-link':
        source=EXTRA/'vinculo.xlsx';rewrite_zip(EXTRA/'origem.xlsx',source,{'xl/externalLinks/externalLink1.xml':'<link/>'})
    elif kind=='xlsx-formula':
        source=EXTRA/'formula.xlsx'
        with zipfile.ZipFile(EXTRA/'origem.xlsx') as archive:data=archive.read('xl/worksheets/sheet1.xml').replace(b'SUM(B3:B4)',b'WEBSERVICE("https://example.invalid")')
        rewrite_zip(EXTRA/'origem.xlsx',source,{'xl/worksheets/sheet1.xml':data})
    elif kind=='ods-link':
        source=EXTRA/'vinculo.ods'
        with zipfile.ZipFile(EXTRA/'origem.ods') as archive:data=archive.read('content.xml').replace(b'of:=SUM([.B3:.B4])',b"of:=['file:///unavailable.ods'#$Sheet1.A1]")
        rewrite_zip(EXTRA/'origem.ods',source,{'content.xml':data})
    elif kind=='odt-image':
        source=EXTRA/'imagem-externa.odt'
        with zipfile.ZipFile(EXTRA/'origem.odt') as archive:data=archive.read('content.xml').replace(b'Pictures/figura.png',b'https://example.invalid/image.png')
        rewrite_zip(EXTRA/'origem.odt',source,{'content.xml':data})
    elif kind=='rtf-field':source=EXTRA/'campo.rtf';source.write_bytes(rb'{\rtf1 {\field{\*\fldinst INCLUDEPICTURE "https://example.invalid/image.png"}}}')
    elif kind=='bad-image':source=EXTRA/'invalida.png';source.write_bytes(b'not a png')
    elif kind=='empty-text':source=EXTRA/'vazio.txt';source.write_bytes(b'')
    else:source=EXTRA/'binario.txt';source.write_bytes(b'hello\x00world')
    code,results=convert([source],'pdf')
    assert code==1 and results[0]['Output'] is None and not source.with_suffix('.pdf').exists(),results


def test_image_batch_isolates_failure_and_preserves_existing_target():
    source=EXTRA/'lote.png';shutil.copy2(EXTRA/'origem.png',source)
    existing=source.with_suffix('.jpg');existing.write_bytes(b'EXISTING DESTINATION')
    invalid=EXTRA/'invalida.png';invalid.write_bytes(b'bad image')
    code,results=convert([invalid,source],'jpg')
    assert code==1 and not results[0]['Success'] and results[1]['Success']
    assert existing.read_bytes()==b'EXISTING DESTINATION' and Path(results[1]['Output']).stem.startswith('lote (')


def test_new_formats_leave_no_partial_files():
    assert not list(EXTRA.rglob('*.partial')) and not list(EXTRA.rglob('*.probe'))
