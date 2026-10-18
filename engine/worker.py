"""Offline document engine. Receives only job snapshots; never publishes final files."""
from __future__ import annotations
import argparse
import contextlib
import io
import json
import logging
import os
from pathlib import Path
import subprocess
import sys
import traceback
import zipfile
import xml.etree.ElementTree as ET
from xml.dom import minidom

# The wire protocol is JSONL; third-party logs are redirected to stderr.
if sys.stdout is not None: sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr is not None: sys.stderr.reconfigure(encoding='utf-8')
PROTOCOL = sys.stdout

class ConversionError(Exception): pass

def emit(kind, **values):
    print(json.dumps(dict(type=kind, **values),ensure_ascii=False),file=PROTOCOL,flush=True)

def progress(message, percent): emit('progress',message=message,percent=percent)
def warning(message): emit('warning',message=message)

def open_pdf(path):
    import pymupdf as fitz
    try:
        doc=fitz.open(path)
        if not doc.is_pdf: raise ConversionError('O arquivo não é um PDF válido.')
        if doc.needs_pass: raise ConversionError('PDF protegido por senha. Abra uma cópia desbloqueada com autorização e tente novamente.')
        if doc.page_count == 0: raise ConversionError('O PDF não contém páginas.')
        if doc.is_repaired: raise ConversionError('O PDF está danificado e precisou de reparo. Salve uma cópia íntegra antes de converter.')
        if doc.is_encrypted and not (doc.permissions & fitz.PDF_PERM_COPY):
            raise ConversionError('O PDF restringe a cópia de conteúdo. Use uma cópia autorizada sem essa restrição.')
        return doc
    except ConversionError: raise
    except Exception as exc: raise ConversionError('Não foi possível abrir o PDF: arquivo inválido ou corrompido.') from exc

def validate_ooxml(path, target, expected_pages=None, check_external=False):
    from PIL import Image
    required='word/document.xml' if target=='docx' else 'ppt/presentation.xml'
    try:
        with open(path,'rb') as stream:
            if stream.read(8)==bytes.fromhex('D0CF11E0A1B11AE1'):
                raise ConversionError('Documento Office protegido por senha ou em formato antigo. Use um DOCX/PPTX desbloqueado.')
        with zipfile.ZipFile(path) as archive:
            infos=archive.infolist()
            if len(infos)>30000 or sum(i.file_size for i in infos)>2*1024**3:
                raise ConversionError('Documento excede o limite de segurança de 2 GB descompactados ou 30 mil entradas.')
            if required not in archive.namelist() or '[Content_Types].xml' not in archive.namelist():
                raise ConversionError('Estrutura DOCX/PPTX incompleta ou extensão incorreta.')
            if archive.testzip(): raise ConversionError('O documento contém dados corrompidos.')
            for info in infos:
                if info.filename.endswith(('.xml','.rels')):
                    content=archive.read(info)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise ConversionError('Documento contém uma declaração XML não permitida.')
                    root=ET.fromstring(content)
                    if check_external and info.filename.endswith('.rels'):
                        for rel in root:
                            if rel.get('TargetMode','').lower()=='external' and not rel.get('Type','').endswith('/hyperlink'):
                                kind=rel.get('Type','').rsplit('/',1)[-1]
                                if kind=='image':
                                    raise ConversionError('O documento contém uma imagem vinculada a outro arquivo, sem incorporação completa. Incorpore essa imagem no Office antes de converter offline.')
                                raise ConversionError('O documento depende de conteúdo externo não incorporado (tipo: '+kind+'). Incorpore esse conteúdo no Office antes de converter offline.')
            if target=='docx':
                from docx import Document
                document=Document(path)
                # Force materialization of text and table structures, not just ZIP creation.
                _=[p.text for p in document.paragraphs]
                _=[[c.text for c in row.cells] for table in document.tables for row in table.rows]
            else:
                from pptx import Presentation
                presentation=Presentation(path)
                if len(presentation.slides)==0: raise ConversionError('A apresentação não contém slides.')
                if expected_pages is not None and len(presentation.slides)!=expected_pages:
                    raise ConversionError('O número de slides não corresponde às páginas do PDF.')
            if not check_external:
                for info in infos:
                    if '/media/' in info.filename and info.filename.lower().endswith(('.png','.jpg','.jpeg')):
                        with Image.open(io.BytesIO(archive.read(info))) as image: image.verify()
    except ConversionError: raise
    except Exception as exc: raise ConversionError('Não foi possível validar o documento Office: arquivo corrompido ou estrutura inválida.') from exc

def prepare_office_input(source,job):
    """Detach an attached Word template in the private job, never in the original.

    DOCX stores its own styles/content. An attached .dot/.dotm is not a missing
    document image: do not fetch it and do not reject an otherwise self-contained
    document. Genuine external content still goes through the strict validator.
    """
    validate_ooxml(source,source.suffix[1:])
    relationship_part='word/_rels/settings.xml.rels'
    settings_part='word/settings.xml'
    template_types={
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate',
        'http://purl.oclc.org/ooxml/officeDocument/relationships/attachedTemplate',
    }
    word_namespaces={
        'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'http://purl.oclc.org/ooxml/wordprocessingml/main',
    }
    relationship_namespaces={
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'http://purl.oclc.org/ooxml/officeDocument/relationships',
    }
    prepared=source
    with zipfile.ZipFile(source) as archive:
        if source.suffix=='.docx' and relationship_part in archive.namelist():
            # DOM retains prefix declarations used by mc:Ignorable. ElementTree
            # reserialization can drop these declarations and damage Office XML.
            relationships=minidom.parseString(archive.read(relationship_part))
            detached=set()
            for relation in list(relationships.documentElement.childNodes):
                if (relation.nodeType==relation.ELEMENT_NODE and relation.localName=='Relationship'
                    and relation.getAttribute('TargetMode').lower()=='external'
                    and relation.getAttribute('Type') in template_types):
                    detached.add(relation.getAttribute('Id'))
                    relationships.documentElement.removeChild(relation)
            if detached:
                if settings_part not in archive.namelist(): raise ConversionError('Configurações do Word ausentes no documento.')
                settings=minidom.parseString(archive.read(settings_part))
                for node in list(settings.getElementsByTagName('*')):
                    if node.namespaceURI not in word_namespaces: continue
                    if node.localName=='attachedTemplate' and any(node.getAttributeNS(ns,'id') in detached for ns in relationship_namespaces):
                        node.parentNode.removeChild(node)
                    elif node.localName=='linkStyles':
                        # Keep the styles saved in the DOCX rather than applying
                        # styles from the now-detached template on load.
                        node.parentNode.removeChild(node)
                replacements={relationship_part:relationships.toxml(encoding='utf-8'),settings_part:settings.toxml(encoding='utf-8')}
                prepared=job/'input-offline.docx'
                if prepared.resolve()==source.resolve(): raise ConversionError('A cópia temporária não pode substituir a origem.')
                with zipfile.ZipFile(prepared,'x') as output:
                    for entry in archive.infolist():
                        data=replacements[entry.filename] if entry.filename in replacements else archive.read(entry)
                        output.writestr(entry,data)
                print('Attached Word template detached from the private conversion snapshot; embedded document content retained.',file=sys.stderr)
    validate_ooxml(prepared,source.suffix[1:],check_external=True)
    return prepared

def office_pdf(source, output, runtime, job):
    source=prepare_office_input(source,job)
    soffice=runtime/'libreoffice'/'program'/'soffice.com'
    if not soffice.exists(): raise ConversionError('LibreOffice incorporado não encontrado. Reinstale o aplicativo.')
    profile=job/'lo-profile'
    (profile/'user').mkdir(parents=True,exist_ok=True)
    # Isolated profile: never launch macros, update links, or share a user's existing LO session.
    (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="FirstRun" oor:op="fuse"><value>false</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Save/Document"><prop oor:name="WarnAlienFormat" oor:op="fuse"><value>false</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Load"><prop oor:name="UpdateDocMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>''',encoding='utf-8')
    outdir=job/'office-output'
    outdir.mkdir(exist_ok=True)
    mode='writer_pdf_Export' if source.suffix=='.docx' else 'impress_pdf_Export'
    options='{"ExportBookmarks":{"type":"boolean","value":"true"},"ExportNotesPages":{"type":"boolean","value":"false"},"ExportHiddenSlides":{"type":"boolean","value":"true"}}'
    args=[str(soffice),f'-env:UserInstallation={profile.as_uri()}','--headless','--nologo','--nodefault','--nofirststartwizard','--norestore',
          '--convert-to',f'pdf:{mode}:{options}','--outdir',str(outdir),str(source)]
    progress('Convertendo com LibreOffice local',25)
    completed=subprocess.run(args,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=1500,
                             creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
    print(completed.stdout.decode('utf-8',errors='replace'),file=sys.stderr)
    print(completed.stderr.decode('utf-8',errors='replace'),file=sys.stderr)
    generated=outdir/(source.stem+'.pdf')
    if completed.returncode or not generated.exists():
        raise ConversionError('O LibreOffice não conseguiu exportar o documento. Verifique se o arquivo está íntegro e sem senha.')
    generated.replace(output)
    return None

def pdf_pptx(source,output):
    import pymupdf as fitz
    from pptx import Presentation
    from pptx.util import Inches
    with open_pdf(source) as pdf:
        deck=Presentation()
        rect=pdf[0].rect
        ratio=rect.width/rect.height
        # PowerPoint supports 1..56 inches per side; extreme PDF ratios use letterboxing.
        ratio=max(1/20,min(20,ratio))
        width,height=(13.333333,13.333333/ratio) if ratio>=1 else (7.5*ratio,7.5)
        width=max(1.0,width); height=max(1.0,height)
        deck.slide_width=Inches(width); deck.slide_height=Inches(height)
        mixed=False
        for index,page in enumerate(pdf):
            progress(f'Renderizando página {index+1} de {len(pdf)}',10+int(70*(index+1)/len(pdf)))
            current=page.rect
            mixed |= abs(current.width/current.height - rect.width/rect.height)>.001
            scale=min(200/72,5000/max(current.width,current.height),(24000000/(current.width*current.height))**.5)
            pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),colorspace=fitz.csRGB,alpha=False)
            image=io.BytesIO(pix.tobytes('png'))
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            scale_emu=min(deck.slide_width/current.width,deck.slide_height/current.height)
            w=int(current.width*scale_emu); h=int(current.height*scale_emu)
            slide.shapes.add_picture(image,(deck.slide_width-w)//2,(deck.slide_height-h)//2,width=w,height=h)
            slide.notes_slide.notes_text_frame.text=f'Página {index+1} do PDF. Dimensão original: {current.width:.2f} x {current.height:.2f} pt. Imagem da página; elementos não editáveis.'
        deck.save(output)
        if mixed: warning('O PDF usa tamanhos de página diferentes; os slides preservam as proporções com margens quando necessário.')
        warning('Cada página foi convertida em uma imagem de alta resolução em um slide. O conteúdo não é editável por elemento.')
        return len(pdf)

def ocr_document(pdf, runtime, job, scanned):
    import pymupdf as fitz
    tessdata=runtime/'ocr'/'tessdata'
    if not all((tessdata/(lang+'.traineddata')).exists() for lang in ['por','eng']):
        raise ConversionError('Dados de OCR local não encontrados. Reinstale o aplicativo com o componente de OCR.')
    # PyMuPDF bundles the Tesseract engine. Only trained language data are external.
    os.environ['TESSDATA_PREFIX']=str(tessdata)
    result=fitz.open()
    for index,page in enumerate(pdf):
        if index not in scanned:
            result.insert_pdf(pdf,from_page=index,to_page=index)
            continue
        progress(f'OCR local: página {index+1} de {len(pdf)}',10+int(40*(index+1)/len(pdf)))
        pix=page.get_pixmap(dpi=220,alpha=False,colorspace=fitz.csRGB)
        ocr=fitz.open(stream=pix.pdfocr_tobytes(language='por+eng',tessdata=str(tessdata)),filetype='pdf')
        if not ocr[0].get_text().strip(): warning(f'Não foi reconhecido texto na página {index+1}; a imagem foi preservada.')
        result.insert_pdf(ocr)
        ocr.close()
    target=job/'ocr.pdf'
    result.save(target)
    result.close()
    warning('OCR local aplicado. Revise palavras, números e tabelas reconhecidos; o layout pode mudar.')
    return target

def pdf_docx(source, output, runtime, job, ocr):
    from pdf2docx import Converter
    converted=source
    ocr_pages=set()
    with open_pdf(source) as pdf:
        pages=len(pdf)
        scanned=[i for i,p in enumerate(pdf) if len(p.get_text().strip())<5 and p.get_images()]
        # PDFs with an existing invisible OCR layer also need the OCR reconstruction mode.
        hidden=[i for i,p in enumerate(pdf) if p.get_text().strip() and p.get_texttrace() and all(s['type']==3 for s in p.get_texttrace())]
        ocr_pages=set(scanned+hidden)
        expected_text=any(p.get_text().strip() for p in pdf)
        if scanned:
            if not ocr: raise ConversionError('O PDF contém páginas digitalizadas. Ative o OCR local nas configurações do Conversor Local e tente novamente.')
            converted=ocr_document(pdf,runtime,job,scanned)
            expected_text=True
    progress('Reconstruindo texto, tabelas e imagens editáveis',55)
    from pdf2docx.page.RawPageFactory import RawPageFactory
    from pdf2docx.page.RawPageFitz import RawPageFitz
    class LocalOcrPage(RawPageFitz):
        def restore(self, **settings):
            # Backend registration is scoped to this isolated one-file worker. Mixed PDFs
            # retain native text/images on native pages and hidden text on OCR pages.
            if self.page_engine.number in ocr_pages: settings=dict(settings,ocr=2)
            return super().restore(**settings)
    original_backend=RawPageFactory.MAP['PYMUPDF']
    RawPageFactory.MAP['PYMUPDF']=LocalOcrPage
    converter=Converter(str(converted))
    try:
        converter.convert(str(output),start=0,end=None,multi_processing=False,ignore_page_error=False)
    finally:
        converter.close()
        RawPageFactory.MAP['PYMUPDF']=original_backend
    from docx import Document
    doc=Document(output)
    texts=[p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]
    if expected_text and not any(x.strip() for x in texts):
        raise ConversionError('A conversão não recuperou texto editável. Nenhum DOCX vazio foi publicado.')
    warning('O PDF não guarda a estrutura original do Word. Revise paginação, colunas, fontes, tabelas e ordem de leitura no DOCX.')
    return pages

def validate_pdf(path, expected=None):
    with open_pdf(path) as document:
        if expected is not None and len(document)!=expected: raise ConversionError('O PDF exportado não contém todos os slides.')
        for page in document:
            if page.rect.is_empty or page.rect.is_infinite: raise ConversionError('Página PDF inválida.')
            # Exercise decoding/content streams on EVERY page before success.
            page.get_pixmap(matrix=__import__('pymupdf').Matrix(.2,.2),alpha=False)
            page.get_text()
        return len(document)

def run(spec):
    source=Path(spec['source']); output=Path(spec['output']); runtime=Path(spec['runtime']); job=Path(spec['job'])
    target=spec['target']
    if (source.suffix,target) not in [('.docx','pdf'),('.pptx','pdf'),('.pdf','docx'),('.pdf','pptx')]:
        raise ConversionError('Operação não suportada.')
    progress('Validando o documento de origem',5)
    expected=None
    if target=='pdf':
        if source.suffix=='.pptx':
            validate_ooxml(source,'pptx',check_external=True)
            from pptx import Presentation
            expected=len(Presentation(source).slides)
        office_pdf(source,output,runtime,job)
    elif target=='pptx': expected=pdf_pptx(source,output)
    else: expected=pdf_docx(source,output,runtime,job,spec.get('ocr',False))
    progress('Reabrindo e validando o arquivo convertido',90)
    if target=='pdf': pages=validate_pdf(output,expected)
    else:
        validate_ooxml(output,target,expected_pages=expected)
        pages=expected or 0
    emit('result',success=True,pages=pages)

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--request',required=True)
    args=parser.parse_args()
    logging.basicConfig(stream=sys.stderr,level=logging.WARNING)
    try:
        spec=json.loads(Path(args.request).read_text(encoding='utf-8-sig'))
        with contextlib.redirect_stdout(sys.stderr): run(spec)
        return 0
    except Exception as exc:
        traceback.print_exc(file=sys.stderr)
        message=str(exc) if isinstance(exc,ConversionError) else 'Falha no motor de conversão. O original foi preservado; consulte o registro local.'
        emit('error',message=message)
        return 1

if __name__=='__main__': sys.exit(main())
